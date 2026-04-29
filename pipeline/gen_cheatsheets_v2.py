"""
HBS Cheat Sheet Generator — v2
Jim Daughtrey | RC MBA 2026
============================================================
WHAT CHANGED FROM v1:
  v1 had one hardcoded builder function per class (build_bgie21,
  build_bgie22, build_fin2_17, etc.). Adding a new week meant
  writing a new function for each class session.

  v2 has a single generalized builder: build_cheatsheet(cls, sections)
  that works for any course and any class. Course colors and folder
  paths are pulled from class_prep_workflow_v2.COURSES. Content
  (case summary, Q&A answers, key frameworks) is passed in as a
  `sections` list by Claude at run time after reading the case PDFs.

USAGE:
  from gen_cheatsheets_v2 import build_cheatsheet

  # cls:      a CLASS_ENTRY dict from class_prep_workflow_v2
  # sections: list of ContentSection dicts (see schema below)
  build_cheatsheet(cls, sections, output_dir="/path/to/folder")

CONTENT SECTION SCHEMA:
  sections = [
      {
          "type":    "info_table",          # 2-col metadata table
          "rows": [
              ("Context",   "text..."),
              ("Key Actors","text..."),
              ("Required Files", "text..."),
          ],
      },
      {
          "type":  "case_summary",           # narrative prose block
          "title": "Case Summary",           # section header text
          "body":  "Multi-paragraph text...",
      },
      {
          "type":  "qa",                     # Q&A block (shaded answer box)
          "title": "Assignment Q&A",         # section header (printed once before first Q)
          "items": [
              {
                  "question": "Q1 text?",
                  "answer":   "A1 text (use \\n for line breaks, • for bullets)",
              },
              {
                  "question": "Q2 text?",
                  "answer":   "A2 text...",
              },
          ],
      },
      {
          "type":  "prose",                  # generic prose section
          "title": "Key Frameworks",
          "items": [
              "Paragraph or bullet text 1",
              "Paragraph or bullet text 2",
          ],
      },
  ]

HISTORY:
  2026-04-07  v2.0  Generalized from per-class builder functions in v1.
"""

import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Import course config from v2 workflow (colors, folder paths)
try:
    from class_prep_workflow_v2 import COURSES, FILE_TYPE_LABELS
except ImportError:
    # Fallback if running standalone — define minimal config
    COURSES = {}
    FILE_TYPE_LABELS = {}


# ─── LOW-LEVEL HELPERS ────────────────────────────────────────────────────────

def hex_to_rgb(h):
    h = h.strip("#")
    return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))

def rgb_color(hex_str):
    r, g, b = hex_to_rgb(hex_str)
    return RGBColor(r, g, b)

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def remove_cell_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'), 'none')
        tcBorders.append(border)
    tcPr.append(tcBorders)

def set_para_font(run, size=10, bold=False, italic=False, color=None, name='Arial'):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = name
    if color:
        run.font.color.rgb = rgb_color(color)

def new_doc():
    """Create a new Document with standard HBS cheat sheet margins and font."""
    doc = Document()
    for section in doc.sections:
        section.top_margin    = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin   = Inches(0.75)
        section.right_margin  = Inches(0.75)
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(4)
    return doc

def spacer(doc, pts=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    p.add_run("").font.size = Pt(pts)


# ─── COMPONENT BUILDERS ───────────────────────────────────────────────────────

def add_header_banner(doc, course, class_num, date_str, time_str, accent_hex):
    """Full-width colored header bar: CLASS X | COURSE | DATE | TIME."""
    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    cell = table.cell(0, 0)
    set_cell_bg(cell, accent_hex)
    remove_cell_borders(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(6)
    p.paragraph_format.left_indent  = Inches(0.1)
    run = p.add_run(f"CLASS {class_num}  |  {course}  |  {date_str}  |  {time_str}")
    set_para_font(run, size=9, bold=True, color="FFFFFF")

def add_title_block(doc, title, subtitle, accent_hex):
    """Large case title + metadata subtitle line."""
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after  = Pt(2)
    run = p_title.add_run(title)
    set_para_font(run, size=14, bold=True, color=accent_hex)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after  = Pt(8)
    run = p_sub.add_run(subtitle)
    set_para_font(run, size=9, italic=True, color="666666")

def add_section_header(doc, text, accent_hex):
    """Colored underlined section heading."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(3)
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'),   'single')
    bottom.set(qn('w:sz'),    '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), accent_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)
    run = p.add_run(text.upper())
    set_para_font(run, size=10, bold=True, color=accent_hex)

def add_info_table(doc, rows_data, accent_hex, light_hex):
    """2-column metadata table: label (colored) | value (light)."""
    table = doc.add_table(rows=len(rows_data), cols=2)
    table.style = 'Table Grid'
    for i, (label, value) in enumerate(rows_data):
        row = table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        set_cell_bg(c0, accent_hex)
        set_cell_bg(c1, light_hex)
        for cell, text, bold, color in [
            (c0, label, True,  "FFFFFF"),
            (c1, value, False, "000000"),
        ]:
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after  = Pt(3)
            p.paragraph_format.left_indent  = Inches(0.07)
            run = p.add_run(text)
            set_para_font(run, size=9, bold=bold, color=color)

def add_prose_block(doc, text, size=10, space_after=6):
    """Plain narrative paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    set_para_font(run, size=size)

def add_qa_block(doc, question, answer, accent_hex, light_hex, q_num):
    """Bold colored question + shaded answer box."""
    # Question line
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(f"Q{q_num}: {question}")
    set_para_font(run, size=10, bold=True, color=accent_hex)

    # Answer in shaded borderless table cell
    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    cell = table.cell(0, 0)
    set_cell_bg(cell, light_hex)
    remove_cell_borders(cell)

    lines = [l.strip() for l in answer.split('\n') if l.strip()]
    first = True
    for line in lines:
        ap = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        ap.paragraph_format.space_before = Pt(3)
        ap.paragraph_format.space_after  = Pt(3)
        ap.paragraph_format.left_indent  = Inches(0.2 if line.startswith('•') else 0.1)
        run = ap.add_run(line)
        set_para_font(run, size=10)


def add_argument_table(doc, question, dimensions, accent_hex, light_hex):
    """
    Side-by-side argument comparison table.
    Layout: 3 columns — Dimension | Argument For | Argument Against
    dimensions: list of {"dimension": str, "for": str, "against": str}
    """
    # Question heading above the table
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(f"Analysis: {question}")
    set_para_font(run, size=10, bold=True, color=accent_hex)

    if not dimensions:
        return

    # Header row + one row per dimension
    table = doc.add_table(rows=1 + len(dimensions), cols=3)
    table.style = 'Table Grid'

    # ── Header row ──────────────────────────────────────────────────────────
    hdr = table.rows[0]
    headers = ["Dimension", "Argument For", "Argument Against"]
    bg_colors = [accent_hex, "1A5C4A", "7B241C"]  # dark navy / dark green / dark red
    for col_idx, (cell, hdr_text, bg) in enumerate(zip(hdr.cells, headers, bg_colors)):
        set_cell_bg(cell, bg)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after  = Pt(3)
        p.paragraph_format.left_indent  = Inches(0.07)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(hdr_text)
        set_para_font(run, size=9, bold=True, color="FFFFFF")

    # ── Data rows ────────────────────────────────────────────────────────────
    # Alternating light-bg rows for readability
    for row_idx, dim in enumerate(dimensions):
        row = table.rows[row_idx + 1]
        row_bg = light_hex if row_idx % 2 == 0 else "FFFFFF"

        dim_cell  = row.cells[0]
        for_cell  = row.cells[1]
        against_cell = row.cells[2]

        # Dimension label (colored background)
        set_cell_bg(dim_cell, accent_hex)
        p = dim_cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after  = Pt(3)
        p.paragraph_format.left_indent  = Inches(0.07)
        run = p.add_run(dim.get("dimension", ""))
        set_para_font(run, size=9, bold=True, color="FFFFFF")

        # For / Against cells
        for cell, text in [(for_cell, dim.get("for", "")),
                           (against_cell, dim.get("against", ""))]:
            set_cell_bg(cell, row_bg)
            lines = [l.strip() for l in text.split('\n') if l.strip()]
            first = True
            for line in lines:
                cp = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                cp.paragraph_format.space_before = Pt(2)
                cp.paragraph_format.space_after  = Pt(2)
                cp.paragraph_format.left_indent  = Inches(0.07)
                run = cp.add_run(line)
                set_para_font(run, size=9)

    spacer(doc, pts=6)


# ─── FILENAME GENERATOR ───────────────────────────────────────────────────────

def cheatsheet_filename(cls):
    """
    Generate the cheat sheet filename from the class entry.
    Convention: "{class_num} {COURSE} {Case Name} Case Analysis Cheatsheet.docx"
    Example:    "23 BGIE AI and Capitalism Case Analysis Cheatsheet.docx"
    """
    import re as _re

    course    = cls.get("course", "")
    class_num = cls["class_num"]
    topic     = cls.get("topic", "Untitled")

    # Strip leading course/class prefixes in all known Canvas formats:
    #   "BGIE | Class 23: "  →  Pattern A
    #   "LCA Class 22 | "    →  Pattern B
    #   "Class 20 | "        →  Pattern C
    _prefix = _re.compile(
        r'^[A-Z]{2,6}\d*\s*\|\s*(?:class\s*\d+\s*[|:–\-]+\s*)?|'   # A
        r'^[A-Z]{2,6}\d*\s+class\s*\d+\s*[|:–\-]+\s*|'              # B
        r'^class\s*\d+\s*[|:–\-]+\s*',                                # C
        _re.I
    )
    case_name = _prefix.sub("", topic).strip().lstrip(":–-|").strip()

    # Also strip the course abbreviation if it starts the case name
    if case_name.upper().startswith(course.upper()):
        case_name = case_name[len(course):].strip().lstrip(":–-|").strip()

    # Remove characters illegal in Windows filenames
    case_name = _re.sub(r'[<>:"/\\|?*]', '', case_name).strip()

    # Cap length
    if len(case_name) > 60:
        case_name = case_name[:60].rsplit(' ', 1)[0].strip()
    if not case_name:
        case_name = topic[:40].strip()

    return f"{class_num} {course} {case_name} Case Analysis Cheatsheet.docx"


# ─── REQUIRED READINGS SUMMARY ────────────────────────────────────────────────

def files_summary(cls):
    """Build a 'Required Files' string for the info table from cls['files']."""
    if not cls["files"]:
        return "No required downloads (build session)"
    parts = []
    for f in cls["files"]:
        label = FILE_TYPE_LABELS.get(f["type"], f["type"].upper())
        parts.append(f"{f['filename']} [{label}]")
    return "\n".join(parts)


# ─── MAIN GENERALIZED BUILDER ─────────────────────────────────────────────────

def build_cheatsheet(cls, sections, output_dir=None):
    """
    Build a cheat sheet .docx for any class session.

    Args:
        cls (dict):       A CLASS_ENTRY dict from class_prep_workflow_v2.
        sections (list):  List of ContentSection dicts (see module docstring).
        output_dir (str): Folder to save the file. If None, uses the course
                          folder from COURSES config.

    Returns:
        str: Path to the saved .docx file.
    """
    course_cfg = COURSES.get(cls["course"], {})
    accent_hex = course_cfg.get("accent", "1F3864")
    light_hex  = course_cfg.get("light",  "EBF0F7")

    # Output path
    if output_dir is None:
        output_dir = course_cfg.get("folder", ".")
    os.makedirs(output_dir, exist_ok=True)
    fname = cheatsheet_filename(cls)
    output_path = os.path.join(output_dir, fname)

    doc = new_doc()

    # ── Header banner ──────────────────────────────────────────────────────
    # Format date: "2026-04-08" → "Wednesday Apr 8, 2026"
    # Use cross-platform day formatting (%-d is Linux only; %#d is Windows only)
    from datetime import datetime
    import platform
    dt = datetime.strptime(cls["date"], "%Y-%m-%d")
    day_fmt = "%#d" if platform.system() == "Windows" else "%-d"
    date_str = dt.strftime(f"%A %b {day_fmt}, %Y")
    add_header_banner(doc, cls["course"], cls["class_num"],
                      date_str, cls["event_time"], accent_hex)
    spacer(doc)

    # ── Title + subtitle ───────────────────────────────────────────────────
    # Build subtitle from files list
    file_refs = []
    for f in cls.get("files", []):
        file_refs.append(f["filename"])
    subtitle_files = " + ".join(file_refs) if file_refs else "No required downloads"
    add_title_block(doc, cls["topic"], subtitle_files, accent_hex)

    # ── Content sections ───────────────────────────────────────────────────
    qa_header_printed = False

    for section in sections:
        stype = section.get("type")

        if stype == "info_table":
            add_info_table(doc, section["rows"], accent_hex, light_hex)
            spacer(doc)

        elif stype in ("case_summary", "prose"):
            add_section_header(doc, section.get("title", "Summary"), accent_hex)
            body = section.get("body") or "\n".join(section.get("items", []))
            add_prose_block(doc, body)

        elif stype == "argument_table":
            add_argument_table(doc,
                               section.get("question", ""),
                               section.get("dimensions", []),
                               accent_hex, light_hex)

        elif stype == "qa":
            if not qa_header_printed:
                add_section_header(doc, section.get("title", "Assignment Q&A"), accent_hex)
                qa_header_printed = True
            for i, item in enumerate(section.get("items", []), start=1):
                add_qa_block(doc,
                             item["question"],
                             item["answer"],
                             accent_hex, light_hex, i)

    # ── If no sections provided, at minimum render the assignment questions ─
    if not sections and cls.get("questions"):
        add_section_header(doc, "Assignment Questions", accent_hex)
        for i, q in enumerate(cls["questions"], start=1):
            add_qa_block(doc, q, "(Answer to be added)", accent_hex, light_hex, i)

    doc.save(output_path)
    print(f"✓  {cls['course']} {cls['class_num']} → {output_path}")
    return output_path


# ─── BATCH RUNNER ─────────────────────────────────────────────────────────────

def build_all(week_classes, sections_map, output_dir_override=None):
    """
    Build cheat sheets for all classes in a week.

    Args:
        week_classes (list):   List of CLASS_ENTRY dicts.
        sections_map (dict):   Keys are (course, class_num) tuples, values are
                               sections lists. E.g.:
                               { ("BGIE", 21): [...sections...], ... }
        output_dir_override:   If set, all files go here instead of course folders.

    Returns:
        list of output file paths.
    """
    paths = []
    for cls in week_classes:
        if cls.get("no_cheatsheet"):
            print(f"  skip  {cls['course']} {cls['class_num']} (build session / no cheat sheet)")
            continue
        key = (cls["course"], cls["class_num"])
        sections = sections_map.get(key, [])
        path = build_cheatsheet(cls, sections, output_dir=output_dir_override)
        paths.append(path)
    return paths


# ─── EXAMPLE / SMOKE TEST ─────────────────────────────────────────────────────

if __name__ == "__main__":
    """Quick smoke test: generate a minimal cheat sheet for a dummy class."""
    import tempfile

    dummy_cls = {
        "date":        "2026-04-13",
        "day":         "Monday",
        "course":      "BGIE",
        "class_num":   23,
        "topic":       "Test Case: Smoke Test",
        "event_name":  "BGIE C",
        "event_time":  "9:10 AM – 10:30 AM",
        "canvas_assignment_id": 0,
        "files": [
            {
                "type":     "case",
                "title":    "Test Case",
                "hbsp_num": "000-000",
                "hbsp_url": "https://example.com",
                "filename": "23 Test - Case.pdf",
                "ext":      "pdf",
            }
        ],
        "questions":      ["What is the core dilemma?", "What would you recommend?"],
        "notes":          "Smoke test only.",
        "no_cheatsheet":  False,
    }

    dummy_sections = [
        {
            "type": "info_table",
            "rows": [
                ("Context",        "Smoke test context."),
                ("Key Actors",     "Alice, Bob"),
                ("Required Files", "23 Test - Case.pdf"),
            ],
        },
        {
            "type":  "case_summary",
            "title": "Case Summary",
            "body":  "This is a smoke test paragraph to verify formatting.",
        },
        {
            "type":  "qa",
            "title": "Assignment Q&A",
            "items": [
                {"question": "What is the core dilemma?", "answer": "The test dilemma.\n• Bullet 1\n• Bullet 2"},
                {"question": "What would you recommend?", "answer": "Recommend proceeding with the test."},
            ],
        },
    ]

    with tempfile.TemporaryDirectory() as tmp:
        path = build_cheatsheet(dummy_cls, dummy_sections, output_dir=tmp)
        print(f"\nSmoke test passed. Output: {path}")
